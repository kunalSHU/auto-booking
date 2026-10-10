#!/usr/bin/env python3
"""
Service price estimation via the OpenAI Batch API.

Typical workflow (nothing is sent to OpenAI until you pass --submit):

  1. Generate payloads only (offline, inspect the .jsonl files):
         python3 llm_pricer_optimized.py Toyota.xlsx

  2. Upload, wait, download, validate and merge into a CSV:
         python3 llm_pricer_optimized.py Toyota.xlsx --submit

     Add --no-wait to submit and exit. Later (even from another terminal):
         python3 llm_pricer_optimized.py Toyota.xlsx --retrieve

  3. Re-merge already-downloaded results without any API calls:
         python3 llm_pricer_optimized.py Toyota.xlsx --merge

Everything is written to ./<input_stem>_batches/ :
  <stem>_batch_N.jsonl      request payloads
  manifest.json             batch ids (so you can resume with --retrieve)
  results_<batch_id>.jsonl  raw OpenAI output
  errors_<batch_id>.jsonl   per-request errors, if any
  <stem>_estimated.csv      input rows with validated estimates filled in
  <stem>_retry.csv          rows that were missing/invalid (use as a new input)
"""
import argparse
import csv
import json
import math
import sys
import time
from pathlib import Path

MODEL = "gpt-4o-mini"
DEFAULT_SERVICES_PER_REQUEST = 50
MAX_REQUESTS_PER_BATCH_FILE = 50_000          # OpenAI Batch API limit
MAX_BYTES_PER_BATCH_FILE = 190 * 1024 * 1024  # OpenAI limit is 200 MB; keep headroom
POLL_SECONDS = 60
TERMINAL_STATUSES = ("completed", "failed", "cancelled", "expired")

ESTIMATION_METHOD = "market_average"
PRICE_FIELDS = ("labor_cost", "parts_cost", "total_price", "price_min", "price_max")

JSON_PROMPT_SYSTEM = """You are an expert automotive service pricing estimator for the Greater Toronto Area (GTA), Ontario, Canada (Toronto, Mississauga, Brampton and surrounding cities).

You will receive ONE JSON object describing a single vehicle and a list of services to price:
{"vehicle_id": 123, "vehicle": "2020 Toyota Camry (SE)", "services": [{"service_id": 1, "service_category": "Brakes", "service_name": "..."}]}

Output ONLY a single valid JSON object. No markdown, code fences, explanations, apologies or conversational text. Never refuse: if exact data is unavailable (e.g. a very new model year), use the most recent comparable model year, or similar vehicles in its class, and give your best professional estimate for the Ontario market.

PRICING RULES:
1. All amounts are in Canadian Dollars (CAD), before HST, reflecting a blend of GTA independent shops and dealerships.
2. labor_cost = realistic book labor hours for THIS specific vehicle and service x a GTA shop rate of $110-$160/hr.
3. parts_cost = parts plus consumables (fluids, filters, gaskets, seals, shop supplies) at typical OEM-equivalent or quality aftermarket prices.
4. Account for the vehicle's specifics: engine size and cylinders, turbo, hybrid/EV, AWD/4WD, trim, and luxury/performance premium. Do not give the same price to a base 4-cylinder and a performance V8.
5. All engine oil services use full synthetic oil, with the correct capacity and filter for the vehicle's engine.
6. If a service does not apply to the vehicle (e.g. a timing belt on a chain-driven engine, or an oil change on an EV), set labor_cost, parts_cost, total_price, price_min and price_max to 0.0.
7. Math: total_price MUST equal labor_cost + parts_cost, rounded to 2 decimals. price_min should be about 10-15% below total_price and price_max about 10-15% above, so price_min <= total_price <= price_max.
8. Use plain numbers only (no currency symbols or commas). Estimate each service independently.
9. Set manual_checked to false and estimation_method to "market_average".

REQUIRED JSON OUTPUT FORMAT:
{
    "estimates": [
        {
            "vehicle_id": 0,
            "service_id": 0,
            "estimation_method": "market_average",
            "manual_checked": false,
            "labor_cost": 0.0,
            "parts_cost": 0.0,
            "total_price": 0.0,
            "price_min": 0.0,
            "price_max": 0.0
        }
    ]
}
Return exactly one estimate for every input service, with the vehicle_id and service_id copied unchanged. Do not add, skip or duplicate any."""


# --------------------------------------------------------------------------- #
# Input helpers
# --------------------------------------------------------------------------- #
def clean(value):
    """Cell -> stripped string; None becomes '' (never the text 'None')."""
    return "" if value is None else str(value).strip()


def to_int(value):
    """Handles '12', 12, and Excel's 12.0. Raises ValueError/TypeError otherwise."""
    return int(float(clean(value)))


def iter_input_rows(file_path):
    """Yield each data row as a dict keyed by header. Supports .csv/.xlsx/.xlsm."""
    suffix = Path(file_path).suffix.lower()
    if suffix == ".csv":
        with open(file_path, encoding="utf-8-sig", newline="") as fh:
            yield from csv.DictReader(fh)
        return

    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            rows = workbook.active.iter_rows(values_only=True)
            headers = [clean(h) for h in next(rows, [])]
            for values in rows:
                values = list(values) + [None] * (len(headers) - len(values))
                yield dict(zip(headers, values))
        finally:
            workbook.close()
        return

    raise ValueError("Input must be a .csv, .xlsx, or .xlsm file")


def row_key(row):
    """(vehicle_id, service_id) or None for blank/invalid rows."""
    try:
        return to_int(row.get("vehicle_id")), to_int(row.get("service_id"))
    except (ValueError, TypeError):
        return None


def vehicle_label(row):
    year = clean(row.get("year"))
    try:
        year = str(to_int(year))
    except (ValueError, TypeError):
        pass
    label = " ".join(p for p in (year, clean(row.get("make")), clean(row.get("model"))) if p)
    trim = clean(row.get("trim"))
    return f"{label} ({trim})" if trim else label


# --------------------------------------------------------------------------- #
# Step 1: generate JSONL payloads (offline)
# --------------------------------------------------------------------------- #
class BatchFileWriter:
    """Writes request lines, rolling to a new file at the request/size limits."""

    def __init__(self, out_dir, stem):
        self.out_dir, self.stem = out_dir, stem
        self.paths, self.handle = [], None
        self.count = self.size = 0

    def write(self, request):
        data = (json.dumps(request) + "\n").encode("utf-8")
        if self.handle and (self.count >= MAX_REQUESTS_PER_BATCH_FILE
                            or self.size + len(data) > MAX_BYTES_PER_BATCH_FILE):
            self.close()
        if not self.handle:
            name = f"{self.stem}_batch_{len(self.paths) + 1}.jsonl"
            self.paths.append(name)
            self.handle = open(self.out_dir / name, "wb")
            self.count = self.size = 0
        self.handle.write(data)
        self.count += 1
        self.size += len(data)

    def close(self):
        if self.handle:
            self.handle.close()
            self.handle = None


def generate_payloads(input_file, out_dir, stem, per_request, max_requests=None):
    out_dir.mkdir(parents=True, exist_ok=True)

    # Group by vehicle so every request covers exactly one vehicle.
    vehicles, seen = {}, set()
    skipped = duplicates = 0
    for row in iter_input_rows(input_file):
        key = row_key(row)
        if key is None:
            skipped += 1
            continue
        if key in seen:
            duplicates += 1
            continue
        seen.add(key)
        vehicle_id, service_id = key
        entry = vehicles.setdefault(vehicle_id, {"vehicle": vehicle_label(row), "services": []})
        entry["services"].append({
            "service_id": service_id,
            "service_category": clean(row.get("service_category")),
            "service_name": clean(row.get("service_name")),
        })

    writer = BatchFileWriter(out_dir, stem)
    request_count = 0
    for vehicle_id, entry in vehicles.items():
        if max_requests is not None and request_count >= max_requests:
            break
        services = entry["services"]
        for chunk_no, start in enumerate(range(0, len(services), per_request)):
            if max_requests is not None and request_count >= max_requests:
                break
            chunk = services[start:start + per_request]
            payload = {"vehicle_id": vehicle_id, "vehicle": entry["vehicle"], "services": chunk}
            writer.write({
                "custom_id": f"v{vehicle_id}_c{chunk_no}",
                "method": "POST",
                "url": "/v1/chat/completions",
                "body": {
                    "model": MODEL,
                    "messages": [
                        {"role": "system", "content": JSON_PROMPT_SYSTEM},
                        {"role": "user", "content": json.dumps(payload, separators=(",", ":"))},
                    ],
                    "temperature": 0.0,
                    "max_completion_tokens": 120 * len(chunk) + 500,
                    "response_format": {"type": "json_object"},
                },
            })
            request_count += 1
    writer.close()

    manifest = {
        "input_file": str(input_file),
        "model": MODEL,
        "batch_files": writer.paths,
        "batches": [],
    }
    save_manifest(out_dir, manifest)
    print(f"Vehicles in input: {len(vehicles)} | services in input: {len(seen)} | "
          f"requests written: {request_count} (<= {per_request} services each)")
    if max_requests is not None:
        print(f"--max-requests {max_requests} applied: test run only, remaining rows will show up in the retry file.")
    if skipped:
        print(f"Skipped {skipped} blank/invalid row(s) (missing vehicle_id or service_id).")
    if duplicates:
        print(f"Skipped {duplicates} duplicate (vehicle_id, service_id) row(s).")
    print(f"Wrote {len(writer.paths)} file(s) to {out_dir}/: {', '.join(writer.paths) or 'none'}")
    return manifest


def save_manifest(out_dir, manifest):
    (out_dir / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")


def load_manifest(out_dir):
    path = out_dir / "manifest.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else None


# --------------------------------------------------------------------------- #
# Step 2: submit / retrieve (OpenAI calls)
# --------------------------------------------------------------------------- #
def make_client():
    from dotenv import load_dotenv
    from openai import OpenAI

    load_dotenv()  # reads .env so OPENAI_API_KEY is available
    return OpenAI()


def submit_batches(client, out_dir, manifest):
    label = Path(manifest["input_file"]).name
    for file_name in manifest["batch_files"]:
        print(f"Uploading '{file_name}'...")
        with open(out_dir / file_name, "rb") as fh:
            uploaded = client.files.create(file=fh, purpose="batch")
        job = client.batches.create(
            input_file_id=uploaded.id,
            endpoint="/v1/chat/completions",
            completion_window="24h",
            metadata={"file": file_name, "description": f"Service estimation - {label}"[:500]},
        )
        manifest["batches"].append({"batch_id": job.id, "file": file_name, "input_file_id": uploaded.id})
        save_manifest(out_dir, manifest)  # saved immediately so --retrieve can resume
        print(f"  Batch ID: {job.id} | status: {job.status}")


def retrieve_batches(client, out_dir, manifest, wait=True):
    pending = {b["batch_id"] for b in manifest["batches"]
               if not (out_dir / f"results_{b['batch_id']}.jsonl").exists()}
    print(f"{len(pending)} batch(es) to retrieve.")

    while pending:
        for batch_id in list(pending):
            batch = client.batches.retrieve(batch_id)
            counts = batch.request_counts
            progress = f"{counts.completed}/{counts.total} done, {counts.failed} failed" if counts else ""
            print(f"Batch {batch_id}: {batch.status} {progress}")

            if batch.status not in TERMINAL_STATUSES:
                continue

            # A batch can finish with BOTH an output file and an error file
            # (or, if everything failed, with only an error file).
            if batch.output_file_id:
                text = client.files.content(batch.output_file_id).text
                (out_dir / f"results_{batch_id}.jsonl").write_text(text, encoding="utf-8")
                print(f"  Saved results_{batch_id}.jsonl")
            if batch.error_file_id:
                text = client.files.content(batch.error_file_id).text
                (out_dir / f"errors_{batch_id}.jsonl").write_text(text, encoding="utf-8")
                print(f"  Saved errors_{batch_id}.jsonl")
            if batch.status != "completed":
                print(f"  Ended with status '{batch.status}'; any rows missing will land in the retry file.")
            if not batch.output_file_id:
                # Write an empty marker so we don't poll this batch forever.
                (out_dir / f"results_{batch_id}.jsonl").write_text("", encoding="utf-8")
            pending.discard(batch_id)

        if pending:
            if not wait:
                print(f"{len(pending)} batch(es) still running. Re-run with --retrieve later.")
                return
            print(f"Waiting {POLL_SECONDS}s... ({len(pending)} still processing)\n")
            time.sleep(POLL_SECONDS)


# --------------------------------------------------------------------------- #
# Step 3: validate + merge
# --------------------------------------------------------------------------- #
def validate_estimate(est):
    """Return (clean_values, reason, repaired). clean_values is None if invalid."""
    try:
        vals = {f: round(float(est[f]), 2) for f in PRICE_FIELDS}
    except (KeyError, TypeError, ValueError):
        return None, "missing_or_non_numeric", False
    if any(not math.isfinite(v) or v < 0 for v in vals.values()):
        return None, "negative_or_non_finite", False

    repaired = False
    expected_total = round(vals["labor_cost"] + vals["parts_cost"], 2)
    if abs(vals["total_price"] - expected_total) > 0.01:
        vals["total_price"] = expected_total  # labor + parts are the primitives
        repaired = True
    if not (vals["price_min"] <= vals["total_price"] <= vals["price_max"]):
        return None, "range_inconsistent", False
    return vals, None, repaired


def parse_results(manifest, out_dir):
    """Collect validated estimates keyed by (vehicle_id, service_id)."""
    estimates, bad_keys, reasons = {}, set(), {}
    failed_requests = repaired = 0

    def reject(reason):
        reasons[reason] = reasons.get(reason, 0) + 1

    for batch in manifest["batches"]:
        path = out_dir / f"results_{batch['batch_id']}.jsonl"
        if not path.exists():
            print(f"Warning: {path.name} not downloaded yet (run --retrieve).")
            continue
        with open(path, encoding="utf-8") as fh:
            for line in fh:
                if not line.strip():
                    continue
                try:
                    obj = json.loads(line)
                    response = obj.get("response") or {}
                    if obj.get("error") or response.get("status_code") != 200:
                        raise ValueError("request failed")
                    request_vehicle = int(obj["custom_id"].split("_")[0][1:])
                    content = response["body"]["choices"][0]["message"]["content"]
                    items = json.loads(content)["estimates"]
                    if not isinstance(items, list):
                        raise ValueError("estimates is not a list")
                except (ValueError, KeyError, IndexError, TypeError, AttributeError):
                    failed_requests += 1
                    continue

                for est in items:
                    try:
                        key = (to_int(est["vehicle_id"]), to_int(est["service_id"]))
                    except (KeyError, TypeError, ValueError):
                        reject("bad_ids")
                        continue
                    if key[0] != request_vehicle:
                        reject("wrong_vehicle_id")
                        continue
                    vals, reason, was_repaired = validate_estimate(est)
                    if vals is None:
                        reject(reason)
                        bad_keys.add(key)
                        continue
                    if key in estimates:
                        reject("duplicate")
                        bad_keys.add(key)  # ambiguous: drop and retry
                        continue
                    estimates[key] = vals
                    repaired += was_repaired

    for key in bad_keys:
        estimates.pop(key, None)
    return estimates, reasons, failed_requests, repaired


def merge_results(manifest, input_file, out_dir, stem):
    estimates, reasons, failed_requests, repaired = parse_results(manifest, out_dir)

    merged_path = out_dir / f"{stem}_estimated.csv"
    retry_path = out_dir / f"{stem}_retry.csv"
    merged = retry = total = 0
    merged_fh = retry_fh = merged_writer = retry_writer = None
    try:
        merged_fh = open(merged_path, "w", encoding="utf-8", newline="")
        for row in iter_input_rows(input_file):
            key = row_key(row)
            if key is None:
                continue  # blank row
            total += 1
            if merged_writer is None:
                fields = list(row.keys())
                merged_writer = csv.DictWriter(merged_fh, fieldnames=fields)
                merged_writer.writeheader()
                retry_fh = open(retry_path, "w", encoding="utf-8", newline="")
                retry_writer = csv.DictWriter(retry_fh, fieldnames=fields)
                retry_writer.writeheader()

            vals = estimates.get(key)
            if vals:
                row.update(vals)
                row["estimation_method"] = ESTIMATION_METHOD
                row["manual_checked"] = "false"
                merged += 1
            else:
                retry_writer.writerow(row)  # original, untouched row
                retry += 1
            merged_writer.writerow(row)
    finally:
        if merged_fh:
            merged_fh.close()
        if retry_fh:
            retry_fh.close()

    if retry == 0 and retry_path.exists():
        retry_path.unlink()

    print("\n--- Merge report ---")
    print(f"Rows in input:            {total}")
    print(f"Estimated and merged:     {merged}")
    print(f"Totals auto-corrected:    {repaired} (total_price reset to labor + parts)")
    print(f"Failed requests:          {failed_requests}")
    for reason, count in sorted(reasons.items()):
        print(f"Rejected ({reason}): {count}")
    print(f"Missing/invalid (retry):  {retry}")
    print(f"Wrote {merged_path}")
    if retry:
        print(f"Wrote {retry_path}  <- run this script on it again to re-estimate those rows")


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def main():
    parser = argparse.ArgumentParser(description="Generate, submit and merge OpenAI service-price estimates.")
    parser.add_argument("file", help="Input CSV or Excel workbook (.csv/.xlsx/.xlsm)")
    parser.add_argument("--submit", action="store_true", help="Upload payloads to OpenAI, wait, then merge")
    parser.add_argument("--retrieve", action="store_true", help="Download results of already-submitted batches, then merge")
    parser.add_argument("--merge", action="store_true", help="Only validate/merge results already downloaded (no API calls)")
    parser.add_argument("--no-wait", action="store_true", help="Do not poll for completion (with --submit/--retrieve)")
    parser.add_argument("--resubmit", action="store_true", help="Allow regenerating/submitting even if batches were already submitted")
    parser.add_argument("--services-per-request", type=int, default=DEFAULT_SERVICES_PER_REQUEST,
                        help=f"Max services per request (default {DEFAULT_SERVICES_PER_REQUEST})")
    parser.add_argument("--max-requests", type=int,
                        help="Only write the first N requests (handy for a small test run)")
    parser.add_argument("--output-dir", help="Output directory (default ./<input_stem>_batches)")
    args = parser.parse_args()

    if args.services_per_request < 1:
        parser.error("--services-per-request must be >= 1")
    if args.max_requests is not None and args.max_requests < 1:
        parser.error("--max-requests must be >= 1")

    input_file = Path(args.file)
    if not input_file.exists():
        sys.exit(f"Input file not found: {input_file}")
    stem = input_file.stem
    out_dir = Path(args.output_dir or f"{stem}_batches")
    existing = load_manifest(out_dir)

    # --- resume / merge modes -------------------------------------------------
    if args.retrieve or args.merge:
        if not existing or not existing.get("batches"):
            sys.exit(f"No submitted batches found in {out_dir}/manifest.json. Run with --submit first.")
        if args.retrieve:
            retrieve_batches(make_client(), out_dir, existing, wait=not args.no_wait)
        merge_results(existing, input_file, out_dir, stem)
        return

    # --- generate (+ optionally submit) ---------------------------------------
    if existing and existing.get("batches"):
        if not args.resubmit:
            sys.exit(f"{out_dir}/ already has submitted batches. Use --retrieve or --merge, "
                     "or pass --resubmit to start a fresh run (the old manifest is backed up).")
        backup = out_dir / f"manifest_old_{int(time.time())}.json"
        (out_dir / "manifest.json").rename(backup)
        print(f"Backed up old manifest to {backup.name}")

    manifest = generate_payloads(input_file, out_dir, stem, args.services_per_request, args.max_requests)
    if not args.submit or not manifest["batch_files"]:
        print("Offline mode: no OpenAI API calls were made. Add --submit to upload these files.")
        return

    client = make_client()
    submit_batches(client, out_dir, manifest)
    if args.no_wait:
        print("Submitted. Later run: python3 llm_pricer_optimized.py "
              f"{input_file} --retrieve")
        return
    retrieve_batches(client, out_dir, manifest, wait=True)
    merge_results(manifest, input_file, out_dir, stem)


if __name__ == "__main__":
    main()